"""Gold Excel ingestion for the 10-gene parkinsonism classifier."""

from __future__ import annotations

import hashlib
import json
from collections import Counter, defaultdict
from collections.abc import Mapping
from dataclasses import dataclass
from pathlib import Path

from deltatorrent.clinical.canonical_features import (
    CANONICAL_PHENOTYPE_FEATURES,
    GENES,
    PEDIGREE_FEATURES,
    CanonicalColumnMapping,
    aao_group_for_years,
    canonicalize_family_history,
    canonicalize_inheritance,
    canonicalize_sex,
    canonicalize_text_features,
    canonicalize_zygosity,
    classify_presence,
    map_column_for_gene,
    normalize_column_name,
    ontology_export,
    parse_aao_years,
)

DEFAULT_DATASET_ROOT = Path("D:/delta_dataset")

_GOLD_FILENAMES: dict[str, str] = {
    "GBA1": "GBA1_PARK.xlsx",
    "LRRK2": "gold-standard.xlsx",
    "PRKN": "PRKN-PARK_20260131.xlsx",
    "GCH1": "gold-standard.xlsx",
    "PINK1": "PINK1_November 2025.xlsx",
    "SNCA": "SNCA (6).xlsx",
    "PARK7": "PARK7_November 2025_SCH (1).xlsx",
    "VPS35": "gold-standard.xlsx",
    "RAB32": "gold-standard.xlsx",
    "ATP13A2": "ATP13A2.xlsx",
}


@dataclass(frozen=True, slots=True)
class FeatureStatus:
    feature: str
    source_column: str
    status: str


@dataclass(frozen=True, slots=True)
class ClinicalRecord:
    record_id: str
    gene: str
    publication_id: str
    family_id: str
    individual_id: str
    sex: str
    aao_years: float | None
    aao_group: str
    aao_missing: bool
    zygosity: str
    inheritance: str
    family_history: str
    phenotype_features: tuple[str, ...]
    pedigree_features: tuple[str, ...]
    feature_statuses: tuple[FeatureStatus, ...]
    source_columns: tuple[tuple[str, str], ...]
    publication_missing: bool

    def stable_identity(self) -> dict[str, object]:
        return {
            "record_id": self.record_id,
            "gene": self.gene,
            "publication_id": self.publication_id,
            "family_id": self.family_id,
            "individual_id": self.individual_id,
        }


def default_gold_files(dataset_root: Path = DEFAULT_DATASET_ROOT) -> dict[str, Path]:
    return {
        gene: dataset_root / gene / "gold" / filename for gene, filename in _GOLD_FILENAMES.items()
    }


def _clean_value(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    if text.lower() in {"-99", "-99.0", "none", "nan", "unknown", "not reported"}:
        return None
    return text


def _clean_identifier(value: object) -> str | None:
    cleaned = _clean_value(value)
    if cleaned is None:
        return None
    lowered = cleaned.lower()
    if lowered.endswith(".0") and lowered[:-2].lstrip("-").isdigit():
        return lowered[:-2]
    return lowered


def _first_existing_index(
    headers: tuple[str, ...],
    mappings: tuple[CanonicalColumnMapping, ...],
    role: str,
    preferred: tuple[str, ...] = (),
) -> int | None:
    by_header = {header: idx for idx, header in enumerate(headers)}
    for header in preferred:
        idx = by_header.get(header)
        if idx is not None and mappings[idx].role == role:
            return idx
    for idx, mapping in enumerate(mappings):
        if mapping.role == role:
            return idx
    return None


def _row_value(row: tuple[object, ...], idx: int | None) -> object | None:
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _publication_id(gene: str, row_idx: int, value: object | None) -> tuple[str, bool]:
    cleaned = _clean_identifier(value)
    if cleaned:
        return cleaned, False
    return f"missing_pmid:{gene}:{row_idx}", True


def _family_id(publication_id: str, row_idx: int, value: object | None) -> str:
    cleaned = _clean_identifier(value)
    if cleaned:
        return cleaned
    return f"missing_family:{publication_id}:{row_idx}"


def _individual_id(row_idx: int, value: object | None) -> str:
    cleaned = _clean_identifier(value)
    if cleaned:
        return cleaned
    return f"missing_individual:{row_idx}"


def _aao_group_from_raw(aao_years: float | None, values: tuple[object | None, ...]) -> str:
    if aao_years is not None:
        return aao_group_for_years(aao_years)
    joined = " ".join(str(v or "").lower() for v in values)
    if "juvenile" in joined or "child" in joined:
        return "JUVENILE_ONSET"
    if "early" in joined:
        return "EARLY_ONSET"
    if "late" in joined:
        return "LATE_ONSET"
    return "AAO_MISSING"


def _stable_record_id(
    *,
    gene: str,
    publication_id: str,
    family_id: str,
    individual_id: str,
    row_idx: int,
    phenotype_features: tuple[str, ...],
) -> str:
    payload = {
        "family_id": family_id,
        "gene": gene,
        "individual_id": individual_id,
        "phenotype_features": list(phenotype_features),
        "publication_id": publication_id,
        "row_idx": row_idx,
    }
    digest = hashlib.sha256(
        json.dumps(payload, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()[:16]
    return f"{gene}:{publication_id}:{family_id}:{individual_id}:{digest}"


def load_gold_records(
    gene_files: Mapping[str, str | Path] | None = None,
    *,
    dataset_root: Path = DEFAULT_DATASET_ROOT,
) -> tuple[ClinicalRecord, ...]:
    """Load one classifier example per clinical row from gold Excel files."""

    files = (
        default_gold_files(dataset_root)
        if gene_files is None
        else {gene: Path(path) for gene, path in gene_files.items()}
    )
    unknown_genes = sorted(set(files) - set(GENES))
    if unknown_genes:
        raise ValueError(f"UNKNOWN_GENE_FILES: {unknown_genes}")

    import openpyxl  # type: ignore[import-untyped]

    records: list[ClinicalRecord] = []
    for gene in GENES:
        path = files.get(gene)
        if path is None:
            continue
        if not path.exists():
            raise FileNotFoundError(path)
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook.worksheets[0]
        raw_headers = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = tuple(
            str(cell or f"col_{idx}").strip().lower() for idx, cell in enumerate(raw_headers)
        )
        normalized_headers = tuple(normalize_column_name(header) for header in headers)
        mappings = tuple(map_column_for_gene(gene, header) for header in normalized_headers)

        publication_idx = _first_existing_index(normalized_headers, mappings, "publication_id")
        family_idx = _first_existing_index(normalized_headers, mappings, "family_id")
        individual_idx = _first_existing_index(normalized_headers, mappings, "individual_id")
        sex_idx = _first_existing_index(normalized_headers, mappings, "sex")
        aao_indices = tuple(idx for idx, mapping in enumerate(mappings) if mapping.role == "aao")
        preferred_aao = ("aao", "aao_movement_disorder", "aao_any", "aao_classify")

        for row_idx, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            row_tuple = tuple(row)
            if not row_tuple or all(cell is None for cell in row_tuple):
                continue

            publication_id, publication_missing = _publication_id(
                gene, row_idx, _row_value(row_tuple, publication_idx)
            )
            family_id = _family_id(publication_id, row_idx, _row_value(row_tuple, family_idx))
            individual_id = _individual_id(row_idx, _row_value(row_tuple, individual_idx))
            sex = canonicalize_sex(_row_value(row_tuple, sex_idx))

            aao_values: list[object | None] = []
            for preferred in preferred_aao:
                if preferred in normalized_headers:
                    aao_values.append(_row_value(row_tuple, normalized_headers.index(preferred)))
            for idx in aao_indices:
                value = _row_value(row_tuple, idx)
                if value not in aao_values:
                    aao_values.append(value)
            aao_years = next(
                (parsed for parsed in map(parse_aao_years, aao_values) if parsed), None
            )
            aao_group = _aao_group_from_raw(aao_years, tuple(aao_values))

            zygosity_values: list[object] = []
            inheritance_values: list[object] = []
            family_history_values: list[object] = []
            phenotype_features: list[str] = []
            pedigree_features: list[str] = []
            feature_statuses: list[FeatureStatus] = []
            source_columns: list[tuple[str, str]] = []

            for col_idx, mapping in enumerate(mappings):
                if col_idx >= len(row_tuple):
                    continue
                header = normalized_headers[col_idx]
                value = row_tuple[col_idx]
                if mapping.role == "phenotype":
                    status = classify_presence(value)
                    for feature in mapping.targets:
                        feature_statuses.append(FeatureStatus(feature, header, status))
                        source_columns.append((feature, header))
                        if status == "positive" and feature not in phenotype_features:
                            phenotype_features.append(feature)
                elif mapping.role == "pedigree":
                    status = classify_presence(value)
                    for feature in mapping.targets:
                        feature_statuses.append(FeatureStatus(feature, header, status))
                        source_columns.append((feature, header))
                        if status == "positive" and feature not in pedigree_features:
                            pedigree_features.append(feature)
                elif mapping.role == "text_features":
                    for feature in canonicalize_text_features(value):
                        if (
                            feature in CANONICAL_PHENOTYPE_FEATURES
                            and feature not in phenotype_features
                        ):
                            phenotype_features.append(feature)
                        elif feature in PEDIGREE_FEATURES and feature not in pedigree_features:
                            pedigree_features.append(feature)
                        feature_statuses.append(FeatureStatus(feature, header, "positive"))
                        source_columns.append((feature, header))
                    family_history_from_text = canonicalize_family_history(value)
                    if family_history_from_text != "unknown":
                        family_history_values.append(family_history_from_text)
                    inheritance_from_text = canonicalize_inheritance(value)
                    if inheritance_from_text != "unknown":
                        inheritance_values.append(inheritance_from_text)
                elif mapping.role == "zygosity":
                    if _clean_value(value):
                        zygosity_values.append(value)
                elif mapping.role == "inheritance":
                    if header.endswith("de_novo") and classify_presence(value) == "positive":
                        inheritance_values.append("de_novo")
                    elif _clean_value(value):
                        inheritance_values.append(value)
                elif mapping.role == "family_history" and _clean_value(value):
                    family_history_values.append(value)
                    for feature in canonicalize_text_features(value):
                        if feature in PEDIGREE_FEATURES and feature not in pedigree_features:
                            pedigree_features.append(feature)
                            feature_statuses.append(FeatureStatus(feature, header, "positive"))
                            source_columns.append((feature, header))

            phenotype_tuple = tuple(
                feature for feature in CANONICAL_PHENOTYPE_FEATURES if feature in phenotype_features
            )
            pedigree_tuple = tuple(
                feature for feature in PEDIGREE_FEATURES if feature in pedigree_features
            )
            zygosity = canonicalize_zygosity(tuple(zygosity_values))
            inheritance = next(
                (
                    parsed
                    for parsed in map(canonicalize_inheritance, inheritance_values)
                    if parsed != "unknown"
                ),
                "unknown",
            )
            family_history = next(
                (
                    parsed
                    for parsed in map(canonicalize_family_history, family_history_values)
                    if parsed != "unknown"
                ),
                "unknown",
            )
            record_id = _stable_record_id(
                gene=gene,
                publication_id=publication_id,
                family_id=family_id,
                individual_id=individual_id,
                row_idx=row_idx,
                phenotype_features=phenotype_tuple,
            )
            records.append(
                ClinicalRecord(
                    record_id=record_id,
                    gene=gene,
                    publication_id=publication_id,
                    family_id=family_id,
                    individual_id=individual_id,
                    sex=sex,
                    aao_years=aao_years,
                    aao_group=aao_group,
                    aao_missing=aao_years is None,
                    zygosity=zygosity,
                    inheritance=inheritance,
                    family_history=family_history,
                    phenotype_features=phenotype_tuple,
                    pedigree_features=pedigree_tuple,
                    feature_statuses=tuple(feature_statuses),
                    source_columns=tuple(sorted(set(source_columns))),
                    publication_missing=publication_missing,
                )
            )

    return tuple(records)


def dataset_summary(records: tuple[ClinicalRecord, ...]) -> dict[str, object]:
    raw_records_per_gene = Counter(record.gene for record in records)
    unique_publications = {record.publication_id for record in records}
    unique_families = {(record.publication_id, record.family_id) for record in records}
    duplicate_keys = Counter(
        (
            record.gene,
            record.publication_id,
            record.family_id,
            record.individual_id,
            record.phenotype_features,
            record.aao_years,
        )
        for record in records
    )
    duplicates = sum(count - 1 for count in duplicate_keys.values() if count > 1)
    return {
        "ontology": ontology_export(),
        "raw_records_per_gene": dict(sorted(raw_records_per_gene.items())),
        "total_records": len(records),
        "unique_publications": len(unique_publications),
        "unique_families": len(unique_families),
        "duplicate_records_estimate": duplicates,
        "missing_publication_records": sum(1 for record in records if record.publication_missing),
    }


def feature_coverage_report(
    records: tuple[ClinicalRecord, ...],
    *,
    split_by_record_id: Mapping[str, str] | None = None,
) -> list[dict[str, object]]:
    statuses: dict[str, Counter[str]] = {
        feature: Counter() for feature in (*CANONICAL_PHENOTYPE_FEATURES, *PEDIGREE_FEATURES)
    }
    source_columns: dict[str, set[str]] = defaultdict(set)
    genes: dict[str, set[str]] = defaultdict(set)
    split_counts: dict[str, Counter[str]] = defaultdict(Counter)

    for record in records:
        positive_features = set(record.phenotype_features) | set(record.pedigree_features)
        for feature in positive_features:
            statuses[feature]["positive"] += 1
            genes[feature].add(record.gene)
            split = split_by_record_id.get(record.record_id) if split_by_record_id else None
            if split:
                split_counts[feature][split] += 1
        for feature_status in record.feature_statuses:
            source_columns[feature_status.feature].add(feature_status.source_column)
            if feature_status.status != "positive":
                statuses[feature_status.feature][feature_status.status] += 1

    report: list[dict[str, object]] = []
    for feature in (*CANONICAL_PHENOTYPE_FEATURES, *PEDIGREE_FEATURES):
        report.append(
            {
                "canonical_feature": feature,
                "source_columns": sorted(source_columns[feature]),
                "records_positive": statuses[feature]["positive"],
                "records_negative": statuses[feature]["negative"],
                "records_unknown": statuses[feature]["unknown"],
                "genes_represented": sorted(genes[feature]),
                "train_count": split_counts[feature]["train"],
                "validation_count": split_counts[feature]["validation"],
                "calibration_count": split_counts[feature]["calibration"],
                "test_count": split_counts[feature]["test"],
            }
        )
    return report
